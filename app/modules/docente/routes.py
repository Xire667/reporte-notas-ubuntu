from flask import render_template, request, redirect, url_for, flash, jsonify, make_response, abort
from flask_login import login_required, current_user
from app import db
from app.models import Usuario, Curso, CursoDocente, CursoAlumno, Nota, NotaActividades, NotaPracticas, NotaParcial, CicloAcademico
from . import docente_bp

def docente_required(f):
    """Decorator para requerir rol de docente"""
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.rol != 'docente':
            flash('No tienes permisos para acceder a esta pÃ¡gina.', 'error')
            return redirect(url_for('main.dashboard'))
        return f(*args, **kwargs)
    return decorated_function

@docente_bp.route('/')
@login_required
@docente_required
def dashboard():
    # Obtener cursos asignados al docente
    cursos_asignados = db.session.query(Curso).join(CursoDocente).filter(
        CursoDocente.docente_id == current_user.id
    ).all()
    # Ciclos presentes en los cursos del docente
    ciclos_ids = {c.ciclo_academico_id for c in cursos_asignados if getattr(c, 'ciclo_academico_id', None)}
    ciclos = CicloAcademico.query.filter(CicloAcademico.id.in_(list(ciclos_ids))).order_by(CicloAcademico.orden).all() if ciclos_ids else []
    
    return render_template('docente/dashboard.html', cursos=cursos_asignados, ciclos=ciclos)

@docente_bp.route('/cursos/<int:curso_id>/alumnos')
@login_required
@docente_required
def ver_alumnos_curso(curso_id):
    # Verificar que el docente tenga asignado este curso
    curso_docente = CursoDocente.query.filter_by(
        curso_id=curso_id, 
        docente_id=current_user.id
    ).first()
    
    if not curso_docente:
        flash('No tienes acceso a este curso.', 'error')
        return redirect(url_for('docente.dashboard'))
    
    # Obtener alumnos del curso
    alumnos = db.session.query(Usuario).join(CursoAlumno).filter(
        CursoAlumno.curso_id == curso_id
    ).all()
    
    curso = Curso.query.get(curso_id)
    
    return render_template('docente/alumnos_curso.html', 
                         alumnos=alumnos, 
                         curso=curso)

@docente_bp.route('/cursos/<int:curso_id>/notas')
@login_required
@docente_required
def gestionar_notas(curso_id):
    """Redirige a la vista de alumnos del curso para seleccionar el alumno especÃ­fico"""
    return redirect(url_for('docente.ver_alumnos_curso', curso_id=curso_id))

@docente_bp.route('/cursos/<int:curso_id>/notas/guardar', methods=['POST'])
@login_required
@docente_required
def guardar_notas(curso_id):
    # Verificar que el docente tenga asignado este curso
    curso_docente = CursoDocente.query.filter_by(
        curso_id=curso_id, 
        docente_id=current_user.id
    ).first()
    
    if not curso_docente:
        return jsonify({'success': False, 'message': 'No tienes acceso a este curso.'})
    
    try:
        # Obtener datos del formulario
        alumno_id = request.form.get('alumno_id')
        comentarios = request.form.get('comentarios', '')
        estado = request.form.get('estado', 'borrador')
        
        # Obtener el curso
        curso = Curso.query.get(curso_id)
        if not curso:
            return jsonify({'success': False, 'message': 'Curso no encontrado.'})
        
        # Validar que se proporcione el alumno_id
        if not alumno_id:
            return jsonify({'success': False, 'message': 'ID de alumno requerido.'})
        
        # Verificar que el alumno estÃ© matriculado en el curso
        matricula = CursoAlumno.query.filter_by(
            curso_id=curso_id,
            alumno_id=alumno_id
        ).first()
        
        if not matricula:
            return jsonify({'success': False, 'message': 'El alumno no estÃ¡ matriculado en este curso.'})
        
        # Recopilar y validar notas de actividades (8 actividades)
        actividades = {}
        for i in range(1, 9):
            valor = request.form.get(f'actividad{i}')
            try:
                actividades[f'actividad{i}'] = float(valor) if valor and valor.strip() != '' else 0.0
            except ValueError:
                return jsonify({'success': False, 'message': f'Error en el formato de la actividad {i}. Debe ser un nÃºmero vÃ¡lido.'})
        
        # Recopilar y validar notas de prÃ¡cticas (4 prÃ¡cticas)
        practicas = {}
        for i in range(1, 5):
            valor = request.form.get(f'practica{i}')
            try:
                practicas[f'practica{i}'] = float(valor) if valor and valor.strip() != '' else 0.0
            except ValueError:
                return jsonify({'success': False, 'message': f'Error en el formato de la prÃ¡ctica {i}. Debe ser un nÃºmero vÃ¡lido.'})
        
        # Recopilar y validar notas de parciales (2 parciales)
        parciales = {}
        for i in range(1, 3):
            valor = request.form.get(f'parcial{i}')
            try:
                parciales[f'parcial{i}'] = float(valor) if valor and valor.strip() != '' else 0.0
            except ValueError:
                return jsonify({'success': False, 'message': f'Error en el formato del parcial {i}. Debe ser un nÃºmero vÃ¡lido.'})
        
        # Validar rango de notas (0-20)
        todas_las_notas = list(actividades.values()) + list(practicas.values()) + list(parciales.values())
        for nota_val in todas_las_notas:
            if nota_val < 0 or nota_val > 20:
                return jsonify({'success': False, 'message': 'Las notas deben estar entre 0 y 20.'})
        
        # Buscar o crear NotaActividades
        nota_actividades = NotaActividades.query.filter_by(
            curso_id=curso_id, 
            alumno_id=alumno_id
        ).first()
        
        if not nota_actividades:
            nota_actividades = NotaActividades(
                curso_id=curso_id,
                alumno_id=alumno_id,
                docente_id=current_user.id
            )
            db.session.add(nota_actividades)
        
        # Actualizar actividades
        for key, value in actividades.items():
            setattr(nota_actividades, key, value)
        nota_actividades.calcular_promedio_actividades()
        
        # Buscar o crear NotaPracticas
        nota_practicas = NotaPracticas.query.filter_by(
            curso_id=curso_id, 
            alumno_id=alumno_id
        ).first()
        
        if not nota_practicas:
            nota_practicas = NotaPracticas(
                curso_id=curso_id,
                alumno_id=alumno_id,
                docente_id=current_user.id
            )
            db.session.add(nota_practicas)
        
        # Actualizar prÃ¡cticas
        for key, value in practicas.items():
            setattr(nota_practicas, key, value)
        nota_practicas.calcular_promedio_practicas()
        
        # Buscar o crear NotaParcial
        nota_parcial = NotaParcial.query.filter_by(
            curso_id=curso_id, 
            alumno_id=alumno_id
        ).first()
        
        if not nota_parcial:
            nota_parcial = NotaParcial(
                curso_id=curso_id,
                alumno_id=alumno_id,
                docente_id=current_user.id
            )
            db.session.add(nota_parcial)
        
        # Actualizar parciales
        for key, value in parciales.items():
            setattr(nota_parcial, key, value)
        nota_parcial.calcular_promedio_parciales()
        
        # Guardar las notas individuales primero
        db.session.commit()
        
        # Buscar o crear la nota principal
        nota = Nota.query.filter_by(
            curso_id=curso_id, 
            alumno_id=alumno_id
        ).first()
        
        if not nota:
            nota = Nota(
                curso_id=curso_id,
                alumno_id=alumno_id,
                docente_id=current_user.id,
                nota_actividades_id=nota_actividades.id,
                nota_practicas_id=nota_practicas.id,
                nota_parcial_id=nota_parcial.id
            )
            db.session.add(nota)
        else:
            # Actualizar referencias si ya existe
            nota.nota_actividades_id = nota_actividades.id
            nota.nota_practicas_id = nota_practicas.id
            nota.nota_parcial_id = nota_parcial.id
        
        nota.comentarios = comentarios
        nota.estado = estado
        
        # Calcular promedio final
        nota.calcular_promedio_final()
        
        db.session.commit()
        
        return jsonify({
            'success': True, 
            'message': 'Notas guardadas correctamente.',
            'promedio_final': nota.promedio_final,
            'promedio_actividades': nota_actividades.promedio_actividades,
            'promedio_practicas': nota_practicas.promedio_practicas,
            'promedio_parciales': nota_parcial.promedio_parciales,
            'estado': nota.estado
        })
        
    except Exception as e:
        db.session.rollback()
        print(f"Error al guardar notas: {e}")  # Para debugging
        return jsonify({'success': False, 'message': f'Error interno del servidor: {str(e)}'})

@docente_bp.route('/cursos/<int:curso_id>/notas/<int:alumno_id>')
@login_required
@docente_required
def ver_nota_alumno(curso_id, alumno_id):
    # Verificar que el docente tenga asignado este curso
    curso_docente = CursoDocente.query.filter_by(
        curso_id=curso_id, 
        docente_id=current_user.id
    ).first()
    
    if not curso_docente:
        flash('No tienes acceso a este curso.', 'error')
        return redirect(url_for('docente.dashboard'))
    
    # Obtener datos del alumno y su nota
    alumno = Usuario.query.get(alumno_id)
    curso = Curso.query.get(curso_id)
    nota = Nota.query.filter_by(
        curso_id=curso_id, 
        alumno_id=alumno_id
    ).first()
    
    # Obtener las notas detalladas
    nota_actividades = NotaActividades.query.filter_by(
        curso_id=curso_id, 
        alumno_id=alumno_id
    ).first()
    
    nota_practicas = NotaPracticas.query.filter_by(
        curso_id=curso_id, 
        alumno_id=alumno_id
    ).first()
    
    nota_parcial = NotaParcial.query.filter_by(
        curso_id=curso_id, 
        alumno_id=alumno_id
    ).first()
    
    return render_template('docente/ver_nota_alumno.html', 
                         alumno=alumno, 
                         curso=curso, 
                         nota=nota,
                         nota_actividades=nota_actividades,
                         nota_practicas=nota_practicas,
                         nota_parcial=nota_parcial)

@docente_bp.route('/cursos/<int:curso_id>/notas/<int:alumno_id>/cambiar-estado', methods=['POST'])
@login_required
@docente_required
def cambiar_estado_nota(curso_id, alumno_id):
    # Verificar que el docente tenga asignado este curso
    curso_docente = CursoDocente.query.filter_by(
        curso_id=curso_id, 
        docente_id=current_user.id
    ).first()
    
    if not curso_docente:
        return jsonify({'success': False, 'message': 'No tienes acceso a este curso.'})
    
    try:
        # Obtener la nota
        nota = Nota.query.filter_by(
            curso_id=curso_id, 
            alumno_id=alumno_id
        ).first()
        
        if not nota:
            return jsonify({'success': False, 'message': 'No se encontrÃ³ la nota.'})
        
        # Cambiar el estado
        if nota.estado == 'borrador':
            nota.estado = 'publicada'
        else:
            nota.estado = 'borrador'
        
        db.session.commit()
        
        return jsonify({
            'success': True, 
            'message': f'Estado cambiado a {nota.estado}.',
            'estado': nota.estado
        })
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Error al cambiar el estado: {str(e)}'})

@docente_bp.route('/reportes')
@login_required
@docente_required
def reportes():
    cursos = db.session.query(Curso).join(CursoDocente).filter(
        CursoDocente.docente_id == current_user.id
    ).order_by(Curso.nombre).all()

    alumnos = db.session.query(Usuario).join(CursoAlumno).join(Curso).join(CursoDocente).filter(
        CursoDocente.docente_id == current_user.id,
        Usuario.rol == 'alumno'
    ).order_by(Usuario.nombre).all()

    alumnos_unicos = []
    vistos = set()
    for a in alumnos:
        if a.id not in vistos:
            alumnos_unicos.append(a)
            vistos.add(a.id)

    return render_template('docente/reportes.html', cursos=cursos, alumnos=alumnos_unicos)

@docente_bp.route('/reportes/curso/<int:curso_id>')
@login_required
@docente_required
def reporte_curso(curso_id):
    from datetime import datetime
    # Seguridad y contexto del docente si aplica
    try:
        curso = Curso.query.get_or_404(curso_id)
    except Exception:
        abort(404)

    # Obtener alumnos del curso (modelo Usuario, rol alumno)
    alumnos = db.session.query(Usuario).join(CursoAlumno).filter(
        CursoAlumno.curso_id == curso_id,
        Usuario.rol == 'alumno'
    ).order_by(Usuario.apellido.asc(), Usuario.nombre.asc()).all()
    datos = []

    for alumno in alumnos:
        na = NotaActividades.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        np = NotaPracticas.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        npa = NotaParcial.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        nota = Nota.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()

        actividades = []
        if na:
            actividades = [na.actividad1, na.actividad2, na.actividad3, na.actividad4, na.actividad5, na.actividad6, na.actividad7, na.actividad8]
        else:
            actividades = []

        practicas = []
        if np:
            practicas = [np.practica1, np.practica2, np.practica3, np.practica4]
        else:
            practicas = []

        parciales = []
        if npa:
            parciales = [npa.parcial1, npa.parcial2]
        else:
            parciales = []

        prom_acts = 0.0
        prom_pracs = 0.0
        prom_parcs = 0.0

        if nota:
            prom_acts = getattr(nota, 'promedio_actividades', 0.0) or 0.0
            prom_pracs = getattr(nota, 'promedio_practicas', 0.0) or 0.0
            prom_parcs = getattr(nota, 'promedio_parciales', 0.0) or 0.0
        
        # Si promedios están en cero, intenta calcular desde tablas de detalle
        if prom_acts == 0.0 and na:
            prom_acts = getattr(na, 'promedio_actividades', 0.0) or 0.0
        if prom_pracs == 0.0 and np:
            prom_pracs = getattr(np, 'promedio_practicas', 0.0) or 0.0
        if prom_parcs == 0.0 and npa:
            prom_parcs = getattr(npa, 'promedio_parciales', 0.0) or 0.0

        promedio_final = round((prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60), 2)
        estado = getattr(nota, 'estado', None) if nota else None

        datos.append({
            'alumno': alumno,
            'actividades': actividades,
            'practicas': practicas,
            'parciales': parciales,
            'promedio_actividades': round(prom_acts, 2) if prom_acts is not None else None,
            'promedio_practicas': round(prom_pracs, 2) if prom_pracs is not None else None,
            'promedio_parciales': round(prom_parcs, 2) if prom_parcs is not None else None,
            'promedio_final': promedio_final,
            'estado': estado
        })

    return render_template('docente/reporte_curso.html', curso=curso, datos=datos)

@docente_bp.route('/reportes/alumno/<int:alumno_id>')
@login_required
@docente_required
def reporte_alumno(alumno_id):

    alumno = Usuario.query.get_or_404(alumno_id)

    cursos = db.session.query(Curso).join(CursoDocente).join(CursoAlumno, CursoAlumno.curso_id == Curso.id).filter(
        CursoDocente.docente_id == current_user.id,
        CursoAlumno.alumno_id == alumno_id
    ).order_by(Curso.nombre).all()

    datos = []
    for curso in cursos:
        nota = Nota.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
        prom_acts = 0.0
        prom_pracs = 0.0
        prom_parcs = 0.0
        prom_final = 0.0
        estado = None

        if nota:
            prom_acts = nota.promedio_actividades or 0.0
            prom_pracs = nota.promedio_practicas or 0.0
            prom_parcs = nota.promedio_parciales or 0.0
            estado = nota.estado

            if prom_acts == 0:
                na = NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
                prom_acts = na.promedio_actividades if na and na.promedio_actividades else 0.0
            if prom_pracs == 0:
                np = NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
                prom_pracs = np.promedio_practicas if np and np.promedio_practicas else 0.0
            if prom_parcs == 0:
                npa = NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
                prom_parcs = npa.promedio_parciales if npa and npa.promedio_parciales else 0.0

            prom_final = (prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60)
        else:
            na = NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
            np = NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
            npa = NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
            prom_acts = na.promedio_actividades if na and na.promedio_actividades else 0.0
            prom_pracs = np.promedio_practicas if np and np.promedio_practicas else 0.0
            prom_parcs = npa.promedio_parciales if npa and npa.promedio_parciales else 0.0
            prom_final = (prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60)

        datos.append({
            'curso': curso,
            'promedio_actividades': prom_acts,
            'promedio_practicas': prom_pracs,
            'promedio_parciales': prom_parcs,
            'promedio_final': prom_final,
            'estado': estado
        })

    return render_template('docente/reporte_alumno.html', alumno=alumno, datos=datos)


@docente_bp.route('/reportes/alumno/<int:alumno_id>/pdf')
@login_required
@docente_required
def reporte_alumno_pdf(alumno_id):
    alumno = Usuario.query.get_or_404(alumno_id)
    cursos = db.session.query(Curso).join(CursoDocente).join(CursoAlumno, CursoAlumno.curso_id == Curso.id).filter(
        CursoDocente.docente_id == current_user.id,
        CursoAlumno.alumno_id == alumno_id
    ).order_by(Curso.nombre).all()
    datos = []
    for curso in cursos:
        nota = Nota.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
        na = NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
        np = NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
        npa = NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=alumno_id).first()
        
        prom_acts = 0.0
        prom_pracs = 0.0
        prom_parcs = 0.0
        prom_final = 0.0
        
        if nota:
            prom_acts = nota.promedio_actividades or 0.0
            prom_pracs = nota.promedio_practicas or 0.0
            prom_parcs = nota.promedio_parciales or 0.0
            if prom_acts == 0 and na:
                prom_acts = na.promedio_actividades if na.promedio_actividades else 0.0
            if prom_pracs == 0 and np:
                prom_pracs = np.promedio_practicas if np.promedio_practicas else 0.0
            if prom_parcs == 0 and npa:
                prom_parcs = npa.promedio_parciales if npa.promedio_parciales else 0.0
            prom_final = (prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60)
        else:
            prom_acts = na.promedio_actividades if na and na.promedio_actividades else 0.0
            prom_pracs = np.promedio_practicas if np and np.promedio_practicas else 0.0
            prom_parcs = npa.promedio_parciales if npa and npa.promedio_parciales else 0.0
            prom_final = (prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60)
        
        # Obtener notas individuales
        actividades = []
        if na:
            actividades = [na.actividad1, na.actividad2, na.actividad3, na.actividad4, na.actividad5, na.actividad6, na.actividad7, na.actividad8]
        else:
            actividades = [0, 0, 0, 0, 0, 0, 0, 0]
        
        practicas = []
        if np:
            practicas = [np.practica1, np.practica2, np.practica3, np.practica4]
        else:
            practicas = [0, 0, 0, 0]
        
        parciales = []
        if npa:
            parciales = [npa.parcial1, npa.parcial2]
        else:
            parciales = [0, 0]
        
        datos.append({
            'curso': curso,
            'actividades': actividades,
            'practicas': practicas,
            'parciales': parciales,
            'promedio_actividades': prom_acts,
            'promedio_practicas': prom_pracs,
            'promedio_parciales': prom_parcs,
            'promedio_final': prom_final
        })
    html = render_template('docente/reporte_alumno_pdf.html', alumno=alumno, datos=datos)
    from xhtml2pdf import pisa
    from io import BytesIO
    result = BytesIO()
    pisa_status = pisa.CreatePDF(html, dest=result)
    if hasattr(pisa_status, 'err') and pisa_status.err:
        flash('Error generando PDF.', 'error')
        return redirect(url_for('docente.reporte_alumno', alumno_id=alumno_id))
    response = make_response(result.getvalue())
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f"attachment; filename=reporte_alumno_{alumno.dni}.pdf"
    return response

@docente_bp.route('/reportes/curso/<int:curso_id>/pdf')
@login_required
@docente_required
def reporte_curso_pdf(curso_id):
    # Construir los mismos datos detallados que en la vista HTML
    try:
        curso = Curso.query.get_or_404(curso_id)
    except Exception:
        abort(404)

    # Obtener alumnos del curso (modelo Usuario, rol alumno)
    alumnos = db.session.query(Usuario).join(CursoAlumno).filter(
        CursoAlumno.curso_id == curso_id,
        Usuario.rol == 'alumno'
    ).order_by(Usuario.apellido.asc(), Usuario.nombre.asc()).all()
    datos = []
    for alumno in alumnos:
        na = NotaActividades.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        np = NotaPracticas.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        npa = NotaParcial.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()
        nota = Nota.query.filter_by(curso_id=curso_id, alumno_id=alumno.id).first()

        actividades = [na.actividad1, na.actividad2, na.actividad3, na.actividad4, na.actividad5, na.actividad6, na.actividad7, na.actividad8] if na else []
        practicas = [np.practica1, np.practica2, np.practica3, np.practica4] if np else []
        parciales = [npa.parcial1, npa.parcial2] if npa else []

        prom_acts = 0.0
        prom_pracs = 0.0
        prom_parcs = 0.0
        if nota:
            prom_acts = getattr(nota, 'promedio_actividades', 0.0) or 0.0
            prom_pracs = getattr(nota, 'promedio_practicas', 0.0) or 0.0
            prom_parcs = getattr(nota, 'promedio_parciales', 0.0) or 0.0
        if prom_acts == 0.0 and na:
            prom_acts = getattr(na, 'promedio_actividades', 0.0) or 0.0
        if prom_pracs == 0.0 and np:
            prom_pracs = getattr(np, 'promedio_practicas', 0.0) or 0.0
        if prom_parcs == 0.0 and npa:
            prom_parcs = getattr(npa, 'promedio_parciales', 0.0) or 0.0

        promedio_final = round((prom_acts * 0.10) + (prom_pracs * 0.30) + (prom_parcs * 0.60), 2)
        estado = getattr(nota, 'estado', None) if nota else None

        datos.append({
            'alumno': alumno,
            'actividades': actividades,
            'practicas': practicas,
            'parciales': parciales,
            'promedio_actividades': round(prom_acts, 2) if prom_acts is not None else None,
            'promedio_practicas': round(prom_pracs, 2) if prom_pracs is not None else None,
            'promedio_parciales': round(prom_parcs, 2) if prom_parcs is not None else None,
            'promedio_final': promedio_final,
            'estado': estado
        })

    html = render_template('docente/reporte_curso_pdf.html', curso=curso, datos=datos)
    from xhtml2pdf import pisa
    from io import BytesIO
    result = BytesIO()
    pisa_status = pisa.CreatePDF(html, dest=result)
    if hasattr(pisa_status, 'err') and pisa_status.err:
        flash('Error generando PDF.', 'error')
        return redirect(url_for('docente.reporte_curso', curso_id=curso_id))

    response = make_response(result.getvalue())
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f"attachment; filename=reporte_curso_{curso.codigo}.pdf"
    return response

@docente_bp.route('/api/cursos_por_ciclo/<int:ciclo_id>')
@login_required
@docente_required
def api_cursos_por_ciclo_docente(ciclo_id):
    cursos = db.session.query(Curso).join(CursoDocente).filter(
        CursoDocente.docente_id == current_user.id,
        Curso.ciclo_academico_id == ciclo_id
    ).order_by(Curso.codigo).all()
    return jsonify([{'id': c.id, 'codigo': c.codigo, 'nombre': c.nombre} for c in cursos])

@docente_bp.route('/notas/descargar_plantilla', methods=['GET'])
@login_required
@docente_required
def descargar_plantilla_notas_docente():
    """Genera una plantilla Excel con alumnos del curso y notas existentes prellenadas"""
    from io import BytesIO
    try:
        import openpyxl
    except ImportError:
        flash('openpyxl no está disponible en el servidor. Contacta al administrador.', 'error')
        return redirect(url_for('docente.dashboard'))

    curso_id = request.args.get('curso_id', type=int)
    if not curso_id:
        flash('Debes seleccionar un curso.', 'error')
        return redirect(url_for('docente.dashboard'))
    curso = Curso.query.get(curso_id)
    if not curso:
        flash('Curso no válido.', 'error')
        return redirect(url_for('docente.dashboard'))
    # Validar que el curso pertenece al docente
    if not CursoDocente.query.filter_by(curso_id=curso.id, docente_id=current_user.id).first():
        flash('No tienes asignado este curso.', 'error')
        return redirect(url_for('docente.dashboard'))

    # Crear workbook y headers
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Notas'
    headers = ['DNI', 'Nombre', 'Apellido'] + [f'Actividad{i}' for i in range(1, 9)] + [f'Practica{i}' for i in range(1, 5)] + ['Parcial1', 'Parcial2']
    ws.append(headers)

    # Alumnos del curso (solo rol alumno activo)
    estudiantes = db.session.query(Usuario).join(CursoAlumno).filter(
        CursoAlumno.curso_id == curso.id,
        Usuario.rol == 'alumno',
        Usuario.activo == True
    ).order_by(Usuario.apellido, Usuario.nombre).all()

    for est in estudiantes:
        act_vals = ['']*8
        prac_vals = ['']*4
        parciales = ['']*2
        na = NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=est.id).first()
        np = NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=est.id).first()
        npa = NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=est.id).first()
        if na:
            act_vals = [na.actividad1 or '', na.actividad2 or '', na.actividad3 or '', na.actividad4 or '', na.actividad5 or '', na.actividad6 or '', na.actividad7 or '', na.actividad8 or '']
        if np:
            prac_vals = [np.practica1 or '', np.practica2 or '', np.practica3 or '', np.practica4 or '']
        if npa:
            parciales = [npa.parcial1 or '', npa.parcial2 or '']
        ws.append([est.dni or '', est.nombre or '', est.apellido or ''] + act_vals + prac_vals + parciales)

    if not estudiantes:
        ws.append(['', '', ''] + ['']*8 + ['']*4 + ['']*2)

    nombre_archivo = f"plantilla_notas_{curso.codigo}.xlsx" if curso and curso.codigo else 'plantilla_notas.xlsx'
    mem = BytesIO()
    wb.save(mem)
    mem.seek(0)
    resp = make_response(mem.getvalue())
    resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = f'attachment; filename={nombre_archivo}'
    return resp
@docente_bp.route('/notas/importar', methods=['POST'])
@login_required
@docente_required
def importar_notas_excel():
    """Importa notas desde Excel para cursos del docente actual"""
    file = request.files.get('archivo_excel')
    if not file:
        flash('No se adjuntó ningún archivo.', 'error')
        return redirect(url_for('docente.dashboard'))

    from werkzeug.utils import secure_filename
    filename = secure_filename(file.filename)
    if not filename.lower().endswith('.xlsx'):
        flash('Formato inválido. Solo se acepta .xlsx.', 'error')
        return redirect(url_for('docente.dashboard'))

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
    except Exception:
        flash('No se pudo leer el archivo Excel.', 'error')
        return redirect(url_for('docente.dashboard'))

    # Curso desde formulario, validar que pertenece al docente
    curso_id = request.form.get('curso_id', type=int)
    if not curso_id:
        flash('Debes seleccionar un curso para importar.', 'error')
        return redirect(url_for('docente.dashboard'))
    curso = Curso.query.get(curso_id)
    if not curso:
        flash('Curso no válido.', 'error')
        return redirect(url_for('docente.dashboard'))
    if not CursoDocente.query.filter_by(curso_id=curso.id, docente_id=current_user.id).first():
        flash('No tienes asignado este curso.', 'error')
        return redirect(url_for('docente.dashboard'))

    estado_form = request.form.get('estado', 'borrador').strip().lower()
    forzar_matricula = request.form.get('forzar_matricula') == 'on'

    # Mapear encabezados
    header_row = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    expected = {
        'DNI': None, 'Nombre': None, 'Apellido': None,
        'Actividad1': None, 'Actividad2': None, 'Actividad3': None, 'Actividad4': None,
        'Actividad5': None, 'Actividad6': None, 'Actividad7': None, 'Actividad8': None,
        'Practica1': None, 'Practica2': None, 'Practica3': None, 'Practica4': None,
        'Parcial1': None, 'Parcial2': None
    }
    idx_map = {}
    for i, h in enumerate(header_row):
        if h in expected:
            idx_map[h] = i

    required_missing = [k for k in ['DNI'] if k not in idx_map]
    if required_missing:
        flash(f'Faltan columnas obligatorias en el Excel: {", ".join(required_missing)}', 'error')
        return redirect(url_for('docente.dashboard'))

    registros_ok = 0
    registros_error = 0
    errores = []

    def to_float(val):
        try:
            if val is None or str(val).strip() == '':
                return 0.0
            return float(val)
        except Exception:
            return 0.0

    for row_idx in range(2, ws.max_row + 1):
        dni = ws.cell(row=row_idx, column=idx_map.get('DNI') + 1).value if 'DNI' in idx_map else None
        if not dni:
            continue

        alumno = Usuario.query.filter_by(dni=str(dni).strip()).first()
        if not alumno:
            registros_error += 1
            errores.append(f'Fila {row_idx}: Alumno con DNI {dni} no encontrado.')
            continue

        # Validar matrícula; si falta, permitir si existen notas o forzar
        matricula = CursoAlumno.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
        if not matricula:
            notas_previas = (
                NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first() or
                NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first() or
                NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first() or
                Nota.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
            )
            if notas_previas:
                pass
            elif forzar_matricula:
                try:
                    nueva = CursoAlumno(curso_id=curso.id, alumno_id=alumno.id)
                    db.session.add(nueva)
                    db.session.flush()
                except Exception as e:
                    registros_error += 1
                    errores.append(f'Fila {row_idx}: No se pudo matricular automáticamente ({str(e)[:90]}).')
                    continue
            else:
                registros_error += 1
                errores.append(f'Fila {row_idx}: Alumno con DNI {dni} no está matriculado en el curso seleccionado.')
                continue

        # Docente es el usuario actual
        docente = current_user

        # Notas
        act_vals = [to_float(ws.cell(row=row_idx, column=idx_map.get(f'Actividad{i}') + 1).value) if f'Actividad{i}' in idx_map else 0.0 for i in range(1, 9)]
        prac_vals = [to_float(ws.cell(row=row_idx, column=idx_map.get(f'Practica{i}') + 1).value) if f'Practica{i}' in idx_map else 0.0 for i in range(1, 5)]
        parcial1 = to_float(ws.cell(row=row_idx, column=idx_map.get('Parcial1') + 1).value) if 'Parcial1' in idx_map else 0.0
        parcial2 = to_float(ws.cell(row=row_idx, column=idx_map.get('Parcial2') + 1).value) if 'Parcial2' in idx_map else 0.0

        na = NotaActividades.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
        if not na:
            na = NotaActividades(curso_id=curso.id, alumno_id=alumno.id, docente_id=docente.id)
            db.session.add(na)
        na.actividad1, na.actividad2, na.actividad3, na.actividad4, na.actividad5, na.actividad6, na.actividad7, na.actividad8 = act_vals
        na.calcular_promedio_actividades()

        np = NotaPracticas.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
        if not np:
            np = NotaPracticas(curso_id=curso.id, alumno_id=alumno.id, docente_id=docente.id)
            db.session.add(np)
        np.practica1, np.practica2, np.practica3, np.practica4 = prac_vals
        np.calcular_promedio_practicas()

        npa = NotaParcial.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
        if not npa:
            npa = NotaParcial(curso_id=curso.id, alumno_id=alumno.id, docente_id=docente.id)
            db.session.add(npa)
        npa.parcial1 = parcial1
        npa.parcial2 = parcial2
        npa.calcular_promedio_parciales()

        nota = Nota.query.filter_by(curso_id=curso.id, alumno_id=alumno.id).first()
        if not nota:
            nota = Nota(curso_id=curso.id, alumno_id=alumno.id, docente_id=docente.id)
            db.session.add(nota)

        nota.nota_actividades_id = na.id
        nota.nota_practicas_id = np.id
        nota.nota_parcial_id = npa.id
        nota.promedio_actividades = na.promedio_actividades or 0.0
        nota.promedio_practicas = np.promedio_practicas or 0.0
        nota.promedio_parciales = npa.promedio_parciales or 0.0
        nota.promedio_final = nota.calcular_promedio_final() or 0.0
        est = estado_form if estado_form in ['borrador', 'publicada'] else 'borrador'
        nota.estado = 'publicada' if est == 'publicada' else 'borrador'

        try:
            db.session.flush()
            registros_ok += 1
        except Exception as e:
            registros_error += 1
            errores.append(f'Fila {row_idx}: Error al guardar ({str(e)[:120]}).')

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash('Error general al guardar las notas importadas.', 'error')
        return redirect(url_for('docente.dashboard'))

    if registros_error:
        flash(f'Importación completada con {registros_ok} filas correctas y {registros_error} con errores.', 'warning')
    else:
        flash(f'Importación completada. {registros_ok} filas procesadas correctamente.', 'success')
    if errores:
        flash('Ejemplos de errores: ' + '; '.join(errores[:5]), 'warning')
    return redirect(url_for('docente.dashboard'))

